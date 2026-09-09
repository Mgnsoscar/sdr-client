"""Spreadsheet run-log export (client side — ui/run_export.py + the row "Export…" buttons).

The agent serves each run's spreadsheet-shaped log (GET /sequence-runs/{id}/log-table); the client
picks an exportable run and turns each unit's tables into an .xlsx. Covered here:

  - run_has_data / recent_runs — which runs are offered (fired-step filter, sequence + plan scoping,
    newest-first, cap 10).
  - _safe_sheet — workbook-legal, unique sheet titles (sanitize illegal chars, 31-char cap, uniquify).
  - build_workbook / save_workbook — one sheet per (name, table); header + rows; a round-trip reload.
  - tables_to_sheets — one unit-named sheet for a single task, "<unit> — <task>" per task otherwise.
  - the row buttons — _SequenceRow and _PlanRow show/hide the Export button on export_ok and clicking
    it calls on_export.

Pure-function tests need no Qt; the row tests instantiate real widgets, so Qt runs offscreen.
"""
import os

import pytest

pytest.importorskip("PyQt6")
pytest.importorskip("openpyxl")
os.environ.setdefault("QT_QPA_PLATFORM", "offscreen")

from PyQt6.QtCore import QEvent
from PyQt6.QtWidgets import QApplication

from api import models as m
from ui import run_export as rx
from ui import plans_tab as pt
from ui import sequences_panel as sp

_app = QApplication.instance() or QApplication([])


@pytest.fixture(autouse=True)
def _flush_deferred_deletes():
    yield
    _app.processEvents()
    _app.sendPostedEvents(None, QEvent.Type.DeferredDelete.value)
    _app.processEvents()


# ── fixtures ─────────────────────────────────────────────────────────────────

def _run(run_id, sequence_id="s1", plan_id="", *, fired=True, when="2030-01-01T00:00:00+00:00",
         state=m.SequenceState.COMPLETED):
    """A SequenceRun whose single step is fired (or not), timed at `when`."""
    step = m.StepFire(anchor="start", offset_s=0.0, action="start", task_name="tx",
                      fire_at=when, fired_actual=(when if fired else None))
    return m.SequenceRun(id=run_id, sequence_id=sequence_id, sequence_name="n", state=state,
                         on_air_at=when, on_air_actual=when, plan_id=plan_id, steps=[step])


def _table(task="tx", columns=("Time", "Power [dBm]"), rows=(("00:00", -60.0), ("00:01", -50.0))):
    return {"task": task, "columns": list(columns), "rows": [list(r) for r in rows]}


# ── run_has_data / recent_runs ────────────────────────────────────────────────

def test_run_has_data_needs_a_fired_step():
    assert rx.run_has_data(_run("r", fired=True)) is True
    assert rx.run_has_data(_run("r", fired=False)) is False
    assert rx.run_has_data(m.SequenceRun(id="r", sequence_id="s1", sequence_name="n",
                                         on_air_at="2030-01-01T00:00:00+00:00")) is False


def test_recent_runs_filters_scopes_orders_and_caps():
    runs = [
        _run("a", "s1", when="2030-01-01T00:00:00+00:00"),
        _run("b", "s1", when="2030-01-01T00:05:00+00:00"),      # newer
        _run("c", "s2", when="2030-01-01T00:09:00+00:00"),      # other sequence → excluded
        _run("d", "s1", fired=False, when="2030-01-01T00:07:00+00:00"),  # no data → excluded
        _run("e", "s1", plan_id="p1", when="2030-01-01T00:03:00+00:00"),
    ]
    # sequence scope only: b (00:05) before e (00:03) before a (00:00); c + d excluded.
    got = rx.recent_runs(runs, "s1")
    assert [r.id for r in got] == ["b", "e", "a"]

    # plan scope: only runs stamped with that plan id.
    assert [r.id for r in rx.recent_runs(runs, "s1", plan_id="p1")] == ["e"]
    assert rx.recent_runs(runs, "s1", plan_id="nope") == []


def test_recent_runs_caps_at_ten_newest_first():
    # 12 runs of the same sequence, ascending time → newest 10 kept, newest first.
    runs = [_run(f"r{i:02d}", "s1", when=f"2030-01-01T00:{i:02d}:00+00:00") for i in range(12)]
    got = rx.recent_runs(runs, "s1")
    assert len(got) == 10
    assert got[0].id == "r11" and got[-1].id == "r02"          # r11..r02 (r01, r00 dropped)


# ── _safe_sheet ────────────────────────────────────────────────────────────────

def test_safe_sheet_sanitizes_truncates_and_uniquifies():
    used: set = set()
    # Illegal chars []:*?/\ become spaces.
    assert rx._safe_sheet("a/b:c*d?e[f]g\\h", used) == "a b c d e f g h"
    # >31 chars is truncated to 31.
    long = rx._safe_sheet("x" * 40, used)
    assert long == "x" * 31 and len(long) == 31
    # A duplicate (case-insensitive) gets a " (2)" suffix, then " (3)", …
    used2: set = set()
    assert rx._safe_sheet("Unit", used2) == "Unit"
    assert rx._safe_sheet("unit", used2) == "unit (2)"
    assert rx._safe_sheet("UNIT", used2) == "UNIT (3)"


# ── build_workbook / save_workbook ───────────────────────────────────────────

def test_build_workbook_one_sheet_per_table_with_header_and_rows():
    wb = rx.build_workbook([("Unit A", _table()), ("Unit B", {"columns": [], "rows": []})])
    assert wb.sheetnames == ["Unit A", "Unit B"]
    ws = wb["Unit A"]
    assert [c.value for c in ws[1]] == ["Time", "Power [dBm]"]  # header row
    assert [c.value for c in ws[2]] == ["00:00", -60.0]         # first data row
    assert [c.value for c in ws[3]] == ["00:01", -50.0]
    assert ws.freeze_panes == "A2"                              # header frozen when there are cols
    assert wb["Unit B"].max_row == 1                            # empty table → just (an empty) header


def test_build_workbook_never_empty_and_sanitizes_names():
    wb = rx.build_workbook([])                                  # no sheets → a single placeholder
    assert wb.sheetnames == ["Sheet"]
    wb2 = rx.build_workbook([("a/b", _table()), ("a/b", _table())])
    assert wb2.sheetnames == ["a b", "a b (2)"]                 # illegal char + uniquified


def test_save_workbook_round_trips(tmp_path):
    import openpyxl
    path = tmp_path / "log.xlsx"
    rx.save_workbook(str(path), [("Unit A", _table())])
    assert path.exists()
    wb = openpyxl.load_workbook(path)
    assert wb.sheetnames == ["Unit A"]
    ws = wb["Unit A"]
    assert [c.value for c in ws[1]] == ["Time", "Power [dBm]"]
    assert [c.value for c in ws[2]] == ["00:00", -60.0]


# ── tables_to_sheets ───────────────────────────────────────────────────────────

def test_tables_to_sheets_single_task_uses_the_unit_name():
    out = rx.tables_to_sheets("Unit A", {"tables": [_table(task="tx")]})
    assert [name for name, _ in out] == ["Unit A"]


def test_tables_to_sheets_multi_task_prefixes_the_unit():
    out = rx.tables_to_sheets("Unit A", {"tables": [_table(task="tx1"), _table(task="tx2")]})
    assert [name for name, _ in out] == ["Unit A — tx1", "Unit A — tx2"]


def test_tables_to_sheets_no_tables_yields_an_empty_unit_sheet():
    out = rx.tables_to_sheets("Unit A", {"tables": []})
    assert out == [("Unit A", {"columns": [], "rows": []})]


def test_tables_to_sheets_localizes_utc_time_to_the_pc_timezone():
    # The agent's UTC "Time" column becomes local Timezone / Date / Time (the PC's tz). Pin the tz
    # so the assertion is deterministic regardless of where the suite runs.
    import time
    prev = os.environ.get("TZ")
    os.environ["TZ"] = "Europe/Oslo"                 # UTC+02:00 on 2026-09-09 (CEST)
    time.tzset()
    try:
        t = _table(columns=("Time", "Power [dBm]"), rows=(("13:55:07.229", -60.0),))
        out = rx.tables_to_sheets("U", {"on_air_at": "2026-09-09T13:55:08+00:00", "tables": [t]})
        _, tbl = out[0]
        assert tbl["columns"] == ["Timezone", "Date", "Time", "Power [dBm]"]
        assert tbl["rows"][0] == ["UTC+02:00", "2026-09-09", "15:55:07.229", -60.0]
    finally:
        if prev is None:
            os.environ.pop("TZ", None)
        else:
            os.environ["TZ"] = prev
        time.tzset()


def test_tables_to_sheets_crossing_midnight_snaps_to_the_right_local_date():
    # A warm-up row a few seconds before an on-air at 00:00:02Z belongs to the previous UTC day; it
    # must not be stamped with the on-air date. Pin UTC so date arithmetic is unambiguous.
    import time
    prev = os.environ.get("TZ")
    os.environ["TZ"] = "UTC"
    time.tzset()
    try:
        t = _table(columns=("Time", "RF on"),
                   rows=(("23:59:58.000", 0), ("00:00:03.000", 1)))
        out = rx.tables_to_sheets("U", {"on_air_at": "2026-09-10T00:00:02+00:00", "tables": [t]})
        _, tbl = out[0]
        assert [r[:3] for r in tbl["rows"]] == [
            ["UTC+00:00", "2026-09-09", "23:59:58.000"],     # snapped back a day
            ["UTC+00:00", "2026-09-10", "00:00:03.000"],
        ]
    finally:
        if prev is None:
            os.environ.pop("TZ", None)
        else:
            os.environ["TZ"] = prev
        time.tzset()


# ── the row "Export…" button ─────────────────────────────────────────────────

def _seq():
    return m.Sequence(id="s1", name="CW tone", steps=[
        m.SequenceStep(anchor="start", offset_s=0, action=m.StepAction.START, task_name="tx"),
        m.SequenceStep(anchor="stop", offset_s=0, action=m.StepAction.STOP, task_name="tx"),
    ])


def test_sequence_row_export_button_shows_and_fires_when_supported():
    seq = _seq()
    clicks = {"n": 0}
    row = sp._SequenceRow(seq, None, on_start=lambda s: None, on_stop=lambda s: None,
                          on_edit=lambda s: None, on_delete=lambda s: None, on_log=lambda s: None,
                          can_run=True, can_edit=False,
                          on_export=lambda s: clicks.__setitem__("n", clicks["n"] + 1),
                          export_ok=True)
    assert row._export.isVisibleTo(row) is True
    row._export.click()
    assert clicks["n"] == 1


def test_sequence_row_export_button_hidden_when_unsupported():
    seq = _seq()
    # No capability → export_ok False → hidden.
    row = sp._SequenceRow(seq, None, on_start=lambda s: None, on_stop=lambda s: None,
                          on_edit=lambda s: None, on_delete=lambda s: None, on_log=lambda s: None,
                          can_run=True, can_edit=False, on_export=lambda s: None, export_ok=False)
    assert row._export.isVisibleTo(row) is False
    # Not a live unit (Library view, can_run=False) → no run controls, export hidden.
    row2 = sp._SequenceRow(seq, None, on_start=lambda s: None, on_stop=lambda s: None,
                           on_edit=lambda s: None, on_delete=lambda s: None, on_log=lambda s: None,
                           can_run=False, can_edit=True, on_export=lambda s: None, export_ok=True)
    assert row2._export.isVisibleTo(row2) is False


def _plan():
    return m.Plan(id="p1", name="Lab plan", items=[
        m.PlanItem(hostname="unit-a", unit_label="Unit A", sequence_id="s1", sequence_name="CW tone")])


def test_plan_row_export_button_shows_and_fires_when_supported():
    plan = _plan()
    clicks = {"n": 0}
    row = pt._PlanRow(plan, [], 0, 0, on_arm=lambda p: None, on_stop=lambda p: None,
                      on_edit=lambda p: None, on_delete=lambda p: None, on_log=lambda p: None,
                      on_export=lambda p: clicks.__setitem__("n", clicks["n"] + 1), export_ok=True)
    assert row._export.isVisibleTo(row) is True
    row._export.click()
    assert clicks["n"] == 1


def test_plan_row_export_button_hidden_when_unsupported():
    plan = _plan()
    row = pt._PlanRow(plan, [], 0, 0, on_arm=lambda p: None, on_stop=lambda p: None,
                      on_edit=lambda p: None, on_delete=lambda p: None, on_log=lambda p: None,
                      on_export=lambda p: None, export_ok=False)
    assert row._export.isVisibleTo(row) is False


def test_run_log_export_method_is_distinct_from_the_yaml_export_method():
    """Regression: each panel's ROW Export button wires ``_on_export_log`` (takes the
    plan/sequence), which must NOT be shadowed by the tab-level YAML ``_on_export`` (no arg).
    A same-name collision made clicking Export on a plan crash: '_on_export() takes 1
    positional argument but 2 were given'."""
    import inspect
    for cls in (pt.PlansTab, sp.SequencesPanel):
        log_params = list(inspect.signature(cls._on_export_log).parameters)
        assert len(log_params) == 2 and log_params[0] == "self", (cls.__name__, log_params)
        yaml_params = list(inspect.signature(cls._on_export).parameters)
        assert yaml_params == ["self"], (cls.__name__, yaml_params)


def test_conditional_row_buttons_are_never_setvisible_while_parentless(monkeypatch):
    """Regression: a button setVisible()'d BEFORE it is added to the layout still has no parent,
    so setVisible(True) briefly makes it a top-level window (a tiny window flashes on Windows)
    until addWidget reparents it. Every conditionally-shown row button (_hold_now/_edit_wb/_export)
    must be parented to the row at creation. Post-construction `.parent()` can't catch this (the
    layout reparents anyway), so spy on setVisible and flag any parentless-visible call."""
    from PyQt6.QtWidgets import QPushButton
    orphaned: list = []
    orig = QPushButton.setVisible

    def spy(self, visible):
        if visible and self.parent() is None:
            orphaned.append(self.text())
        orig(self, visible)

    monkeypatch.setattr(QPushButton, "setVisible", spy)
    sp._SequenceRow(_seq(), None, on_start=lambda s: None, on_stop=lambda s: None,
                    on_edit=lambda s: None, on_delete=lambda s: None, on_log=lambda s: None,
                    can_run=True, can_edit=False, on_hold_now=lambda s: None,
                    on_edit_wb=lambda s: None, on_export=lambda s: None, export_ok=True)
    pt._PlanRow(_plan(), [], 0, 0, on_arm=lambda p: None, on_stop=lambda p: None,
                on_edit=lambda p: None, on_delete=lambda p: None, on_log=lambda p: None,
                can_ff=True, on_hold_now=lambda p: None, on_export=lambda p: None, export_ok=True)
    assert orphaned == [], f"buttons setVisible()'d while parentless (top-level flash): {orphaned}"
