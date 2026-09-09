"""
Export a ran sequence/plan's log as a spreadsheet.

The agent serves each run's spreadsheet-shaped log (GET /sequence-runs/{id}/log-table): one table
per duration task, each a per-change time-series (Time + every power quantity + realized SDR gain /
attenuation + each parameter + derived readouts; a row only where a value changed). This turns those
tables into an .xlsx — one sheet per unit (a multi-unit plan run → several sheets).

`recent_runs` / `run_has_data` (pure) pick the exportable runs; `build_workbook` / `save_workbook`
(pure, openpyxl) write the file; `RunExportDialog` is the row's "Export log…" entry point: it lists
the last ≤10 runs of a sequence/plan and exports the chosen one.
"""
from __future__ import annotations

import re
from typing import Dict, List, Optional, Tuple

from PyQt6.QtCore import Qt
from PyQt6.QtWidgets import (
    QDialog, QDialogButtonBox, QFileDialog, QLabel, QListWidget, QListWidgetItem,
    QMessageBox, QVBoxLayout,
)

from api import models as m
from .qt_adapter import DataHub
from .theme import Palette
from .widgets import fit_dialog_to_screen


def run_has_data(run: m.SequenceRun) -> bool:
    """True if the run fired at least one step (something to tabulate)."""
    return any(getattr(s, "fired_actual", None) for s in (run.steps or []))


def _run_time(run: m.SequenceRun) -> str:
    return (run.on_air_actual or run.started_actual or run.on_air_at or run.created_at or "")


def recent_runs(runs: List[m.SequenceRun], sequence_id: str, plan_id: str = "",
                limit: int = 10) -> List[m.SequenceRun]:
    """The most recent (≤limit) runs of a sequence (optionally scoped to a plan) that carry data,
    newest first."""
    out = [r for r in runs
           if r.sequence_id == sequence_id
           and (not plan_id or getattr(r, "plan_id", "") == plan_id)
           and run_has_data(r)]
    out.sort(key=_run_time, reverse=True)
    return out[:limit]


def _safe_sheet(name: str, used: set) -> str:
    """A workbook-legal, unique sheet title (≤31 chars, no []:*?/\\)."""
    s = re.sub(r"[\[\]:*?/\\]", " ", (name or "Sheet")).strip()[:31] or "Sheet"
    base, n = s, 2
    while s.lower() in used:
        suffix = f" ({n})"
        s = (base[:31 - len(suffix)] + suffix)
        n += 1
    used.add(s.lower())
    return s


def build_workbook(sheets: List[Tuple[str, dict]]):
    """An openpyxl Workbook, one sheet per (name, table) — table is {columns, rows}. Empty tables
    still get a header sheet. Raises RuntimeError if openpyxl isn't available."""
    try:
        import openpyxl
        from openpyxl.styles import Alignment, Font, PatternFill
        from openpyxl.utils import get_column_letter
    except ImportError as exc:                       # pragma: no cover - packaging guard
        raise RuntimeError("openpyxl is required to export a spreadsheet") from exc

    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    hdr_fill = PatternFill("solid", fgColor="1F6FA5")
    hdr_font = Font(bold=True, color="FFFFFF")
    used: set = set()
    for name, table in (sheets or [("Sheet", {"columns": [], "rows": []})]):
        ws = wb.create_sheet(_safe_sheet(name, used))
        cols = list(table.get("columns") or [])
        ws.append(cols)
        for cix in range(1, len(cols) + 1):
            c = ws.cell(row=1, column=cix)
            c.fill, c.font = hdr_fill, hdr_font
            c.alignment = Alignment(horizontal="center")
        for row in table.get("rows") or []:
            ws.append(list(row))
        for cix, name_ in enumerate(cols, 1):
            ws.column_dimensions[get_column_letter(cix)].width = max(10, min(34, len(str(name_)) + 3))
        if cols:
            ws.freeze_panes = "A2"
    if not wb.sheetnames:                            # never leave a workbook with no sheet
        wb.create_sheet("Sheet")
    return wb


def save_workbook(path: str, sheets: List[Tuple[str, dict]]) -> None:
    build_workbook(sheets).save(path)


def tables_to_sheets(unit_label: str, log_table: dict) -> List[Tuple[str, dict]]:
    """Split one unit's /log-table payload into (sheet_name, table) pairs — one sheet named by the
    unit when it ran a single duration task, else '<unit> — <task>' per task."""
    tables = log_table.get("tables") or []
    if len(tables) <= 1:
        t = tables[0] if tables else {"columns": [], "rows": []}
        return [(unit_label, t)]
    return [(f"{unit_label} — {t.get('task', '')}", t) for t in tables]


class RunExportDialog(QDialog):
    """Pick one of the last ≤10 runs of a sequence/plan and export it to an .xlsx (one sheet per
    unit). `targets` is [(hostname, unit_label)] — one for a sequence, one-per-unit for a plan;
    the run list comes from the first target and the workbook gathers every target's table."""

    def __init__(self, hub: DataHub, targets: List[Tuple[str, str]], sequence_id: str,
                 title: str, *, plan_id: str = "", parent=None):
        super().__init__(parent)
        self.hub = hub
        self.targets = targets
        self.sequence_id = sequence_id
        self.plan_id = plan_id
        self._runs: List[m.SequenceRun] = []
        self.setWindowTitle(f"Export log — {title}")
        self.setMinimumSize(460, 320)
        self._build(title)
        fit_dialog_to_screen(self, 560, 460)
        self.hub.task_done.connect(self._on_task_done)
        self.finished.connect(lambda _=0: self._disconnect())
        self._load_runs()

    def _build(self, title: str) -> None:
        outer = QVBoxLayout(self)
        outer.setContentsMargins(16, 14, 16, 12)
        outer.setSpacing(9)
        note = QLabel("Pick a run to export. Each row of the spreadsheet is one moment the state "
                      "changed — a launch, a tune, RF on/off — with every parameter in its own "
                      "column. A multi-unit plan writes one sheet per unit.")
        note.setWordWrap(True)
        note.setStyleSheet(f"font-size: 11px; color: {Palette.TEXT_MUTED};")
        outer.addWidget(note)

        self._list = QListWidget()
        self._list.setStyleSheet("font-size: 12px;")
        self._list.itemDoubleClicked.connect(lambda _=None: self._export())
        outer.addWidget(self._list, stretch=1)

        self._status = QLabel("loading runs…")
        self._status.setStyleSheet(f"font-size: 11px; color: {Palette.TEXT_FAINT};")
        outer.addWidget(self._status)

        self._buttons = QDialogButtonBox()
        self._export_btn = self._buttons.addButton("Export…", QDialogButtonBox.ButtonRole.AcceptRole)
        self._export_btn.setObjectName("primary")
        self._export_btn.setEnabled(False)
        self._buttons.addButton(QDialogButtonBox.StandardButton.Cancel)
        self._buttons.accepted.connect(self._export)
        self._buttons.rejected.connect(self.reject)
        outer.addWidget(self._buttons)

    # ── Runs ────────────────────────────────────────────────────────────────
    def _load_runs(self) -> None:
        host = self.targets[0][0] if self.targets else ""
        self.hub.run_async(f"export_runs:{host}",
                           lambda h=host: self.hub.fleet.get(h).list_sequence_runs())

    def _on_task_done(self, label: str, result) -> None:
        if label.startswith("export_runs:"):
            runs = result if isinstance(result, list) else []
            self._runs = recent_runs(runs, self.sequence_id, self.plan_id)
            self._list.clear()
            for r in self._runs:
                state = r.state.value if hasattr(r.state, "value") else str(r.state)
                when = _run_time(r).replace("T", " ")[:19]
                item = QListWidgetItem(f"{when}   ·   {state}")
                item.setData(Qt.ItemDataRole.UserRole, r.id)
                self._list.addItem(item)
            if self._runs:
                self._list.setCurrentRow(0)
                self._export_btn.setEnabled(True)
                self._set_status(f"{len(self._runs)} run(s) — newest first")
            else:
                self._set_status("no runs with data to export yet", faint=True)
        elif label.startswith("export_table:"):
            self._on_tables(result)

    # ── Export ──────────────────────────────────────────────────────────────
    def _export(self) -> None:
        item = self._list.currentItem()
        if item is None:
            return
        self._run_id = item.data(Qt.ItemDataRole.UserRole)
        self._set_status("building spreadsheet…")
        self._export_btn.setEnabled(False)
        # One /log-table per target (unit); assemble the sheets when all return.
        self._pending = {h: lbl for h, lbl in self.targets}
        self._sheets: List[Tuple[str, dict]] = []
        for host, label in self.targets:
            self.hub.run_async(
                f"export_table:{host}",
                lambda h=host, lbl=label: (lbl, self.hub.fleet.get(h).sequence_run_log_table(self._run_id)))

    def _on_tables(self, result) -> None:
        try:
            label, payload = result
        except Exception:                            # noqa: BLE001
            self._set_status("could not build the table", error=True)
            self._export_btn.setEnabled(True)
            return
        self._sheets.extend(tables_to_sheets(label, payload))
        # Remove this unit from pending (keyed by label — unique per target here).
        for h, lbl in list(self._pending.items()):
            if lbl == label:
                self._pending.pop(h, None)
                break
        if self._pending:
            return                                    # wait for the rest
        default = f"{self.windowTitle().split('— ', 1)[-1]}.xlsx".replace("/", "-")
        path, _ = QFileDialog.getSaveFileName(self, "Save run log", default,
                                              "Excel workbook (*.xlsx)")
        if not path:
            self._set_status("export cancelled")
            self._export_btn.setEnabled(True)
            return
        if not path.lower().endswith(".xlsx"):
            path += ".xlsx"
        try:
            save_workbook(path, self._sheets)
        except Exception as exc:                      # noqa: BLE001
            QMessageBox.warning(self, "Export failed", str(exc))
            self._set_status("export failed", error=True)
            self._export_btn.setEnabled(True)
            return
        self.accept()

    def _set_status(self, text: str, error: bool = False, faint: bool = False) -> None:
        color = Palette.CRASH if error else (Palette.TEXT_FAINT if faint else Palette.TEXT_MUTED)
        self._status.setText(text)
        self._status.setStyleSheet(f"font-size: 11px; color: {color};")

    def _disconnect(self) -> None:
        try:
            self.hub.task_done.disconnect(self._on_task_done)
        except (TypeError, RuntimeError):
            pass
